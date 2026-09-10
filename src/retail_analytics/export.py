"""Export the tested analytics reports as a Power BI-ready Excel workbook."""

from __future__ import annotations

import argparse
import json
import os
import tempfile
from collections.abc import Mapping
from pathlib import Path
from typing import Final

import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo

from .analytics import QueryResult, run_all_reports
from .audit import audit_reports
from .loader import load_curated_dataset

DEFAULT_OUTPUT: Final[Path] = Path("outputs/dashboard/retail_analytics.xlsx")

REPORT_SHEETS: Final[dict[str, str]] = {
    "kpi_summary": "KPI Summary",
    "monthly_sales": "Monthly Sales",
    "product_performance": "Products",
    "country_performance": "Countries",
    "customer_monthly": "Customers Monthly",
    "data_quality_summary": "Data Quality",
}

TABLE_NAMES: Final[dict[str, str]] = {
    report_name: "tbl" + "".join(part.title() for part in report_name.split("_"))
    for report_name in REPORT_SHEETS
}

MONEY_COLUMNS: Final[frozenset[str]] = frozenset(
    {
        "gross_sales",
        "net_recorded_value",
        "average_order_value",
        "cancellation_value",
        "new_customer_sales",
        "repeat_customer_sales",
        "sale_value",
        "adjustment_value",
    }
)

DATE_COLUMNS: Final[frozenset[str]] = frozenset(
    {
        "transaction_month",
        "order_month",
        "first_transaction_date",
        "last_transaction_date",
    }
)

TIMESTAMP_COLUMNS: Final[frozenset[str]] = frozenset(
    {"earliest_timestamp", "latest_timestamp"}
)

HEADER_FILL: Final[str] = "17365D"
HEADER_FONT: Final[str] = "FFFFFF"
TABLE_STYLE: Final[str] = "TableStyleMedium2"


def _dataframe(result: QueryResult) -> pd.DataFrame:
    """Convert exact SQL rows into a labelled table without recalculating."""
    return pd.DataFrame.from_records(
        result.rows,
        columns=result.columns,
        coerce_float=False,
    )


def _column_number_format(column_name: str) -> str | None:
    if column_name in MONEY_COLUMNS:
        return "#,##0.000"
    if column_name in DATE_COLUMNS:
        return "yyyy-mm-dd"
    if column_name in TIMESTAMP_COLUMNS:
        return "yyyy-mm-dd hh:mm:ss"
    if column_name.endswith("_pct"):
        return "0.00"
    return None


def _format_worksheet(
    worksheet: object,
    report_name: str,
    dataframe: pd.DataFrame,
) -> None:
    """Apply restrained table formatting and explicit data formats."""
    worksheet.freeze_panes = "A2"
    worksheet.sheet_view.showGridLines = False

    for cell in worksheet[1]:
        cell.fill = PatternFill("solid", fgColor=HEADER_FILL)
        cell.font = Font(name="Arial", size=10, bold=True, color=HEADER_FONT)
        cell.alignment = Alignment(horizontal="center", vertical="center")

    worksheet.row_dimensions[1].height = 24

    for column_index, column_name in enumerate(dataframe.columns, start=1):
        number_format = _column_number_format(str(column_name))
        if number_format:
            for cell in worksheet.iter_cols(
                min_col=column_index,
                max_col=column_index,
                min_row=2,
                max_row=max(worksheet.max_row, 2),
            ):
                for item in cell:
                    item.number_format = number_format

        values = [str(column_name)] + [
            "" if value is None else str(value)
            for value in dataframe.iloc[:, column_index - 1].tolist()
        ]
        width = min(max(len(value) for value in values) + 2, 42)
        worksheet.column_dimensions[
            worksheet.cell(row=1, column=column_index).column_letter
        ].width = max(width, 12)

    last_column = worksheet.cell(
        row=1,
        column=max(len(dataframe.columns), 1),
    ).column_letter
    last_row = max(len(dataframe.index) + 1, 1)
    table = Table(
        displayName=TABLE_NAMES[report_name],
        ref=f"A1:{last_column}{last_row}",
    )
    table.tableStyleInfo = TableStyleInfo(
        name=TABLE_STYLE,
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )
    worksheet.add_table(table)


def _write_workbook(
    reports: Mapping[str, QueryResult],
    output_path: Path,
) -> None:
    """Write all report tables atomically so a failed export cannot corrupt output."""
    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    workbook.remove(workbook.active)
    workbook.properties.title = "Retail Sales and Customer Analytics"
    workbook.properties.subject = "Dashboard-ready tables generated from DuckDB SQL"

    for report_name, sheet_name in REPORT_SHEETS.items():
        dataframe = _dataframe(reports[report_name])
        worksheet = workbook.create_sheet(sheet_name)

        rows = [list(dataframe.columns)] + dataframe.values.tolist()
        for row in rows:
            worksheet.append(row)

        _format_worksheet(worksheet, report_name, dataframe)

    temporary_path: Path | None = None
    try:
        with tempfile.NamedTemporaryFile(
            dir=output_path.parent,
            prefix=f".{output_path.stem}-",
            suffix=".xlsx",
            delete=False,
        ) as temporary_file:
            temporary_path = Path(temporary_file.name)

        workbook.save(temporary_path)

        # Reopening detects malformed workbook structures before replacement.
        verification = load_workbook(temporary_path, read_only=False, data_only=True)
        verification.close()
        os.replace(temporary_path, output_path)
    finally:
        workbook.close()
        if temporary_path is not None and temporary_path.exists():
            temporary_path.unlink()


def export_dashboard_workbook(
    source: str | Path,
    output_path: str | Path = DEFAULT_OUTPUT,
    *,
    s3_client: object | None = None,
) -> dict[str, int]:
    """Run all reports and export one named Excel table per report."""
    with load_curated_dataset(source, s3_client=s3_client) as dataset:
        reports = run_all_reports(dataset.connection)

    audit_reports(reports)
    destination = Path(output_path).expanduser().resolve()
    _write_workbook(reports, destination)
    return {report_name: len(result.rows) for report_name, result in reports.items()}


def main() -> int:
    parser = argparse.ArgumentParser(
        description="Export retail analytics tables for Power BI."
    )
    parser.add_argument("source", help="Local Parquet file/directory or S3 URI")
    parser.add_argument(
        "--output",
        default=str(DEFAULT_OUTPUT),
        help=f"Excel destination (default: {DEFAULT_OUTPUT})",
    )
    arguments = parser.parse_args()

    output_path = Path(arguments.output).expanduser().resolve()
    report_rows = export_dashboard_workbook(arguments.source, output_path)
    print(
        json.dumps(
            {"workbook": str(output_path), "report_rows": report_rows},
            indent=2,
        )
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
