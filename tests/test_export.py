from __future__ import annotations

import tempfile
import unittest
from pathlib import Path

import pyarrow as pa
import pyarrow.parquet as pq
from openpyxl import load_workbook

from retail_analytics.export import (
    REPORT_SHEETS,
    TABLE_NAMES,
    export_dashboard_workbook,
)
from retail_analytics.loader import EXPECTED_ARROW_SCHEMA
from tests.test_analytics import TEST_ROWS


class ExportTests(unittest.TestCase):
    def setUp(self) -> None:
        self.temporary_directory = tempfile.TemporaryDirectory()
        self.root = Path(self.temporary_directory.name)
        self.source = self.root / "curated.parquet"
        table = pa.Table.from_pylist(TEST_ROWS, schema=EXPECTED_ARROW_SCHEMA)
        pq.write_table(table, self.source)

    def tearDown(self) -> None:
        self.temporary_directory.cleanup()

    def test_exports_six_named_power_bi_tables(self) -> None:
        output = self.root / "nested" / "retail_analytics.xlsx"

        report_rows = export_dashboard_workbook(self.source, output)

        self.assertTrue(output.is_file())
        self.assertEqual(tuple(report_rows), tuple(REPORT_SHEETS))

        workbook = load_workbook(output, data_only=True)
        try:
            self.assertEqual(workbook.sheetnames, list(REPORT_SHEETS.values()))
            for report_name, sheet_name in REPORT_SHEETS.items():
                worksheet = workbook[sheet_name]
                self.assertIn(TABLE_NAMES[report_name], worksheet.tables)
                self.assertEqual(worksheet.freeze_panes, "A2")
                self.assertFalse(worksheet.sheet_view.showGridLines)
        finally:
            workbook.close()

    def test_preserves_typed_kpi_values_and_formats(self) -> None:
        output = self.root / "retail_analytics.xlsx"
        export_dashboard_workbook(self.source, output)

        workbook = load_workbook(output, data_only=True)
        try:
            worksheet = workbook["KPI Summary"]
            headers = {
                cell.value: cell.column for cell in worksheet[1]
            }

            self.assertEqual(worksheet.cell(2, headers["curated_rows"]).value, 9)
            self.assertEqual(worksheet.cell(2, headers["gross_sales"]).value, 75)
            self.assertEqual(
                worksheet.cell(2, headers["net_recorded_value"]).value,
                57.5,
            )
            self.assertEqual(
                worksheet.cell(2, headers["gross_sales"]).number_format,
                "#,##0.000",
            )
            self.assertEqual(
                worksheet.cell(2, headers["first_transaction_date"]).number_format,
                "yyyy-mm-dd",
            )
        finally:
            workbook.close()


if __name__ == "__main__":
    unittest.main()
